# Copyright (c) Microsoft. All rights reserved.
# Licensed under the MIT license. See LICENSE file in the project root for full license information.

"""Read and structurally validate a HOBL scenario-authoring workbook."""

from __future__ import annotations

import argparse
from io import BytesIO
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


SCENARIO_SHEET = "Scenario"
ACTIONS_SHEET = "Actions"
SCENARIO_NAME_PATTERN = re.compile(r"^[a-z][a-z0-9_]*$")
SUPPORTED_PLATFORMS = {"windows": "Windows", "macos": "macOS"}
ACTION_COLUMNS = (
    "sequence",
    "scenario",
    "actions",
    "inter-action_delay_seconds",
    "expected_result",
    "enabled",
    "notes",
)


class WorkbookValidationError(ValueError):
    """Raised when a workbook does not satisfy the authoring contract."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_positive_integer(value: Any, field: str) -> int:
    if isinstance(value, bool):
        raise WorkbookValidationError(f"{field} must be a positive integer.")
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise WorkbookValidationError(f"{field} must be a positive integer.") from None
    if str(number) != _text(value) and not isinstance(value, int):
        try:
            if float(value) != number:
                raise WorkbookValidationError(f"{field} must be a positive integer.")
        except (TypeError, ValueError):
            raise WorkbookValidationError(f"{field} must be a positive integer.") from None
    if number <= 0:
        raise WorkbookValidationError(f"{field} must be a positive integer.")
    return number


def _parse_nonnegative_number(value: str, field: str) -> int | float:
    try:
        number = float(value)
    except ValueError:
        raise WorkbookValidationError(f"{field} must be a nonnegative number.") from None
    if not math.isfinite(number) or number < 0:
        raise WorkbookValidationError(f"{field} must be a nonnegative number.")
    return int(number) if number.is_integer() else number


def _parse_nonnegative_integer(value: Any, field: str) -> int:
    if isinstance(value, bool):
        raise WorkbookValidationError(f"{field} must be a nonnegative integer.")
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise WorkbookValidationError(f"{field} must be a nonnegative integer.") from None
    if not math.isfinite(number) or number < 0 or not number.is_integer():
        raise WorkbookValidationError(f"{field} must be a nonnegative integer.")
    return int(number)


def _parse_enabled(value: Any, field: str) -> bool:
    if value is None or _text(value) == "":
        return True
    if isinstance(value, bool):
        return value
    normalized = _text(value).lower()
    if normalized in {"1", "true", "yes", "y", "enabled"}:
        return True
    if normalized in {"0", "false", "no", "n", "disabled"}:
        return False
    raise WorkbookValidationError(
        f"{field} must be TRUE/FALSE, YES/NO, or 1/0."
    )


def _read_scenario_sheet(sheet) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for row in sheet.iter_rows(values_only=True):
        key = _text(row[0] if row else None).lower()
        if not key or key == "field":
            continue
        value = row[1] if len(row) > 1 else None
        if key in values:
            raise WorkbookValidationError(
                f"{SCENARIO_SHEET} sheet contains duplicate field {key!r}."
            )
        values[key] = value

    required = ("name", "platform", "description", "typing_delay_ms")
    missing = [field for field in required if not _text(values.get(field))]
    if missing:
        raise WorkbookValidationError(
            f"{SCENARIO_SHEET} sheet is missing required field(s): {', '.join(missing)}."
        )

    name = _text(values["name"])
    if not SCENARIO_NAME_PATTERN.fullmatch(name):
        raise WorkbookValidationError(
            f"Scenario name must use lowercase snake_case: {name!r}."
        )

    platform_key = _text(values["platform"]).lower()
    if platform_key not in SUPPORTED_PLATFORMS:
        raise WorkbookValidationError(
            f"Unsupported platform {_text(values['platform'])!r}; use Windows or macOS."
        )
    if platform_key == "macos" and not name.startswith("mac_"):
        raise WorkbookValidationError(
            f"macOS scenario names must start with 'mac_': {name!r}."
        )

    typing_delay = _parse_nonnegative_integer(
        values["typing_delay_ms"],
        f"{SCENARIO_SHEET}.typing_delay_ms",
    )

    return {
        "name": name,
        "platform": SUPPORTED_PLATFORMS[platform_key],
        "description": _text(values["description"]),
        "author": _text(values.get("author")),
        "capture_device": _text(values.get("capture_device")),
        "typing_delay_ms": typing_delay,
    }


def _read_actions_sheet(sheet) -> list[dict[str, Any]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration:
        raise WorkbookValidationError(f"{ACTIONS_SHEET} sheet is empty.") from None

    headers = [_text(value).lower() for value in header_row]
    missing_columns = [column for column in ACTION_COLUMNS if column not in headers]
    if missing_columns:
        raise WorkbookValidationError(
            f"{ACTIONS_SHEET} sheet is missing required column(s): "
            f"{', '.join(missing_columns)}."
        )
    column_indexes = {column: headers.index(column) for column in ACTION_COLUMNS}

    actions: list[dict[str, Any]] = []
    seen_sequences: set[int] = set()
    current_scenario = ""
    for source_row, row in enumerate(rows, start=2):
        if not any(_text(value) for value in row):
            continue

        def cell(column: str) -> Any:
            index = column_indexes[column]
            return row[index] if index < len(row) else None

        sequence = _parse_positive_integer(
            cell("sequence"),
            f"{ACTIONS_SHEET} row {source_row} sequence",
        )
        if sequence in seen_sequences:
            raise WorkbookValidationError(
                f"{ACTIONS_SHEET} sheet contains duplicate sequence {sequence}."
            )
        seen_sequences.add(sequence)

        scenario = _text(cell("scenario"))
        if scenario:
            current_scenario = scenario
        elif not current_scenario:
            raise WorkbookValidationError(
                f"{ACTIONS_SHEET} row {source_row} scenario is required because "
                "there is no preceding scenario to inherit."
            )

        instruction = _text(cell("actions"))
        if not instruction:
            raise WorkbookValidationError(
                f"{ACTIONS_SHEET} row {source_row} actions is required."
            )

        delay = _parse_nonnegative_number(
            _text(cell("inter-action_delay_seconds")),
            f"{ACTIONS_SHEET} row {source_row} inter-action_delay_seconds",
        )

        actions.append(
            {
                "source_sheet": ACTIONS_SHEET,
                "source_row": source_row,
                "sequence": sequence,
                "scenario": current_scenario,
                "instruction": instruction,
                "inter_action_delay_seconds": delay,
                "expected_result": _text(cell("expected_result")),
                "enabled": _parse_enabled(
                    cell("enabled"), f"{ACTIONS_SHEET} row {source_row} enabled"
                ),
                "notes": _text(cell("notes")),
            }
        )

    if not actions:
        raise WorkbookValidationError(
            f"{ACTIONS_SHEET} sheet must contain at least one action."
        )

    sequences = [action["sequence"] for action in actions]
    expected_sequences = list(range(1, len(actions) + 1))
    if sequences != expected_sequences:
        raise WorkbookValidationError(
            f"Sequence must be consecutive and start at 1; found {sequences}."
        )
    return actions


def read_workbook(path: str | Path) -> dict[str, Any]:
    """Return normalized workbook contents without modifying the workbook."""

    workbook_path = Path(path).expanduser().resolve()
    if not workbook_path.is_file():
        raise WorkbookValidationError(f"Workbook not found: {workbook_path}")
    if workbook_path.suffix.lower() != ".xlsx":
        raise WorkbookValidationError("Workbook must use the .xlsx file extension.")
    if load_workbook is None:
        raise WorkbookValidationError(
            "openpyxl is not installed. Install the repository requirements."
        )

    try:
        workbook_data = BytesIO(workbook_path.read_bytes())
    except OSError as error:
        raise WorkbookValidationError(
            f"Unable to read workbook {workbook_path}: {error}"
        ) from error
    try:
        workbook = load_workbook(workbook_data, read_only=True, data_only=True)
    except Exception as error:
        workbook_data.close()
        raise WorkbookValidationError(
            f"Unable to open workbook {workbook_path}: {error}"
        ) from error

    try:
        missing_sheets = [
            name
            for name in (SCENARIO_SHEET, ACTIONS_SHEET)
            if name not in workbook.sheetnames
        ]
        if missing_sheets:
            raise WorkbookValidationError(
                f"Workbook is missing required sheet(s): {', '.join(missing_sheets)}."
            )
        scenario = _read_scenario_sheet(workbook[SCENARIO_SHEET])
        actions = _read_actions_sheet(workbook[ACTIONS_SHEET])
    finally:
        workbook.close()
        workbook_data.close()

    return {
        "workbook": str(workbook_path),
        "scenario": scenario,
        "actions": actions,
        "summary": {
            "action_count": len(actions),
            "enabled_action_count": sum(action["enabled"] for action in actions),
        },
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Read and validate a HOBL scenario-authoring workbook."
    )
    parser.add_argument("workbook", help="Path to the .xlsx workbook.")
    parser.add_argument(
        "--compact",
        action="store_true",
        help="Write compact JSON instead of indented JSON.",
    )
    args = parser.parse_args(argv)

    try:
        result = read_workbook(args.workbook)
    except WorkbookValidationError as error:
        print(f" ERROR - {error}", file=sys.stderr)
        return 1

    if args.compact:
        print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
    else:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
